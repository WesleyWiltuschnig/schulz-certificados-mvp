"""Leitura e saneamento da planilha .xlsx (RF02 a RF06, regras H01 a H13)."""
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

BLOCOS_ESPERADOS = 2
TAMANHO_MAXIMO = 5 * 1024 * 1024

CAMPOS = {
    "instrumento": {"instrumento", "codigo_do_instrumento", "instrumento_codigo"},
    "denominacao": {"denominacao", "descricao_do_instrumento"},
    "plano_medicao": {"plano_de_medicao", "measurement_plan"},
    "data_medicao": {"data_da_medicao", "data", "date"},
    "temperatura": {"temperatura_ambiente", "temperatura"},
    "padrao_relatorio": {"padrao_do_relatorio", "padrao_utilizado", "mmc"},
    "metrologista": {"metrologista"},
    "lea_padrao": {"lea_padrao", "lea_geral"},
}
OBRIGATORIOS = {
    "instrumento": "Instrumento",
    "denominacao": "Denominação",
    "plano_medicao": "Plano de medição",
    "data_medicao": "Data da medição",
    "temperatura": "Temperatura ambiente",
    "padrao_relatorio": "Padrão do relatório",
}
COL_REFERENCIA = {"valor_padrao", "valor_de_referencia", "referencia"}
COL_LEA = {"lea", "limite_de_erro_admissivel"}
COL_DESVIO = {"desvio", "desvio_informado"}


class IngestaoInvalida(Exception):
    def __init__(self, erros):
        super().__init__("; ".join(e["mensagem"] for e in erros))
        self.erros = erros


def _erro(codigo, mensagem):
    return {"codigo": codigo, "mensagem": mensagem}


def normalizar_chave(valor):
    s = unicodedata.normalize("NFKD", str(valor))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def limpar_texto(valor):
    """H01: remove controle, colapsa espacos e apara as extremidades."""
    if valor is None:
        return ""
    s = "".join(c for c in str(valor) if unicodedata.category(c)[0] != "C" or c in " \t")
    return re.sub(r"\s+", " ", s).strip()


def codigo_normalizado(valor):
    """Compara codigos de padrao ignorando pontuacao (500.011 == 500011)."""
    return re.sub(r"[^A-Za-z0-9]", "", limpar_texto(valor)).upper()


def para_decimal(valor):
    """Converte celula em Decimal sem passar por ponto flutuante binario (H13)."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        raise ValueError("valor booleano")
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and (valor != valor or valor in (float("inf"), float("-inf"))):
            raise ValueError("valor nao finito")
        return Decimal(repr(valor))
    s = limpar_texto(valor).replace(",", ".")
    if s == "":
        return None
    if not re.fullmatch(r"[+-]?\d+(\.\d+)?", s):
        raise ValueError(f"valor nao numerico: {s!r}")
    return Decimal(s)


def extrair_temperatura(valor):
    """H06: separa valor numerico e simbolo (20,4° -> 20.4)."""
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return Decimal(repr(valor))
    s = limpar_texto(valor)
    m = re.fullmatch(r"([+-]?\d+(?:[.,]\d+)?)\s*(?:°|º|graus)?\s*(?:C|c)?", s)
    if not m:
        raise ValueError(f"temperatura nao numerica: {s!r}")
    return Decimal(m.group(1).replace(",", "."))


def interpretar_data(valor):
    """H03/H04/H05: retorna data ISO 8601; sentinela e invalida."""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    s = limpar_texto(valor)
    if re.fullmatch(r"0+[/.\-]0+[/.\-]0+", s):
        raise ValueError("data-sentinela (00/00/0000)")
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, formato).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"data invalida: {s!r}")


def ler_planilha(conteudo, padrao_esperado=None, blocos_esperados=BLOCOS_ESPERADOS):
    """Le o .xlsx e devolve dados saneados; levanta IngestaoInvalida com todos os erros."""
    if len(conteudo) > TAMANHO_MAXIMO:
        raise IngestaoInvalida([_erro("ARQ", "Arquivo acima do tamanho máximo de 5 MB.")])
    if conteudo[:2] != b"PK" or not zipfile.is_zipfile(io.BytesIO(conteudo)):
        raise IngestaoInvalida([_erro("ARQ", "Arquivo não é uma planilha .xlsx válida.")])
    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
        linhas = list(wb.worksheets[0].iter_rows(values_only=True))
    except Exception:
        raise IngestaoInvalida([_erro("ARQ", "Não foi possível ler a planilha.")])

    erros, avisos = [], []
    bruto, idx_cab = {}, None
    for i, linha in enumerate(linhas):
        primeira = linha[0] if linha else None
        if primeira is not None and normalizar_chave(primeira) == "ponto":
            idx_cab = i
            break
        if primeira is not None and len(linha) > 1 and linha[1] not in (None, ""):
            bruto[normalizar_chave(primeira)] = linha[1]
    if idx_cab is None:
        raise IngestaoInvalida([_erro("EST", "Cabeçalho da tabela de pontos (coluna 'Ponto') não encontrado.")])

    meta = {}
    for canonico, apelidos in CAMPOS.items():
        for chave, valor in bruto.items():
            if chave in apelidos:
                meta[canonico] = valor
                break

    for chave, rotulo in OBRIGATORIOS.items():
        if limpar_texto(meta.get(chave)) == "":
            erros.append(_erro("H10", f"Campo obrigatório ausente: {rotulo}."))

    dados = {"avisos": avisos}
    for chave in ("instrumento", "denominacao", "plano_medicao", "padrao_relatorio", "metrologista"):
        dados[chave] = limpar_texto(meta.get(chave))

    if dados["plano_medicao"] and re.search(r"[_\-\s]TESTE$", dados["plano_medicao"], re.I):
        erros.append(_erro("RN10", f"Plano de medição '{dados['plano_medicao']}' indica dado de teste; ingestão recusada."))

    if limpar_texto(meta.get("data_medicao")) != "" or isinstance(meta.get("data_medicao"), (date, datetime)):
        try:
            dados["data_medicao"] = interpretar_data(meta["data_medicao"])
        except ValueError as e:
            erros.append(_erro("H03", f"Data da medição inválida: {e}."))
    if limpar_texto(meta.get("temperatura")) != "":
        try:
            dados["temperatura"] = extrair_temperatura(meta["temperatura"])
        except ValueError as e:
            erros.append(_erro("RN14", f"Temperatura ambiente inválida: {e}."))

    if padrao_esperado is not None and dados["padrao_relatorio"]:
        if codigo_normalizado(dados["padrao_relatorio"]) != codigo_normalizado(padrao_esperado):
            erros.append(_erro("RN03", f"Padrão do relatório ({dados['padrao_relatorio']}) diverge do padrão selecionado ({padrao_esperado})."))

    lea_padrao = None
    if limpar_texto(meta.get("lea_padrao")) != "":
        try:
            lea_padrao = para_decimal(meta["lea_padrao"])
        except ValueError as e:
            erros.append(_erro("LEA", f"LEA geral inválido: {e}."))

    cab = linhas[idx_cab]
    col = {"ponto": None, "ref": None, "lea": None, "desvio": None}
    col_leituras = {}
    for j, c in enumerate(cab):
        n = normalizar_chave(c) if c is not None else ""
        if n == "ponto":
            col["ponto"] = j
        elif n in COL_REFERENCIA:
            col["ref"] = j
        elif n in COL_LEA:
            col["lea"] = j
        elif n in COL_DESVIO:
            col["desvio"] = j
        elif re.fullmatch(r"leitura_(\d+)", n):
            col_leituras[int(n.split("_")[1])] = j
    if col["ref"] is None:
        erros.append(_erro("EST", "Coluna do valor padrão (Valor_Padrao) não encontrada."))
    if len(col_leituras) != blocos_esperados:
        erros.append(_erro("RN01", f"Esperadas {blocos_esperados} leituras por ponto; encontradas {len(col_leituras)} colunas de leitura."))

    pontos, nomes = [], set()
    if col["ref"] is not None and len(col_leituras) == blocos_esperados:
        ordem_cols = [col_leituras[k] for k in sorted(col_leituras)]

        def celula(linha, j):
            return linha[j] if j is not None and j < len(linha) else None

        for n_linha, linha in enumerate(linhas[idx_cab + 1:], start=idx_cab + 2):
            nome = limpar_texto(celula(linha, col["ponto"]))
            if nome == "":
                if any(v not in (None, "") for v in linha):
                    avisos.append(f"Linha {n_linha} suprimida: ponto sem identificação.")
                continue
            if nome in nomes:
                erros.append(_erro("EST", f"Ponto duplicado: {nome}."))
                continue
            nomes.add(nome)
            try:
                ref = para_decimal(celula(linha, col["ref"]))
                if ref is None:
                    raise ValueError("valor padrão ausente")
                leituras = []
                for j in ordem_cols:
                    v = para_decimal(celula(linha, j))
                    if v is None:
                        raise ValueError("leitura ausente")
                    leituras.append(v)
                lea = para_decimal(celula(linha, col["lea"])) if col["lea"] is not None else None
                if lea is None:
                    lea = lea_padrao
                if lea is None:
                    raise ValueError("LEA não informado")
                if lea <= 0:
                    raise ValueError("LEA deve ser positivo")
                desvio = para_decimal(celula(linha, col["desvio"])) if col["desvio"] is not None else None
            except ValueError as e:
                erros.append(_erro("EST", f"Ponto '{nome}' (linha {n_linha}): {e}."))
                continue
            pontos.append({"nome": nome, "referencia": ref, "leituras": leituras, "lea": lea, "desvio_informado": desvio})
        if not pontos and not erros:
            erros.append(_erro("EST", "Nenhum ponto de ensaio encontrado."))

    if erros:
        raise IngestaoInvalida(erros)
    dados["pontos"] = pontos
    dados["blocos"] = blocos_esperados
    return dados
